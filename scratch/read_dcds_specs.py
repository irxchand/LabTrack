import pypdf
import docx
import sys
import pptx
import os

sys.stdout.reconfigure(encoding='utf-8')

print("=== 1. DCDSL_Evaluation_Plan.pdf ===")
pdf_path = r"docs\academic\DCDSL_Evaluation_Plan.pdf"
if os.path.exists(pdf_path):
    reader = pypdf.PdfReader(pdf_path)
    print(f"Total pages: {len(reader.pages)}")
    for i, page in enumerate(reader.pages):
        print(f"--- PAGE {i+1} ---")
        print(page.extract_text()[:2000])
else:
    print("PDF not found!")

print("\n=== 2. DCDSL PROJECT REPORT format.docx ===")
docx_path = r"presentations\DCDS\DCDSL PROJECT REPORT format.docx"
if os.path.exists(docx_path):
    doc = docx.Document(docx_path)
    for p in doc.paragraphs:
        if p.text.strip():
            print(p.text)
    for table in doc.tables:
        print("--- TABLE ---")
        for row in table.rows:
            print([cell.text.strip() for cell in row.cells])
else:
    print("DOCX not found!")

print("\n=== 3. DCDS_ESE_Project_Presentation_Template.pptx ===")
pptx_path = r"presentations\DCDS\DCDS_ESE_Project_Presentation_Template.pptx"
if os.path.exists(pptx_path):
    prs = pptx.Presentation(pptx_path)
    print(f"Total slides: {len(prs.slides)}")
    for i, slide in enumerate(prs.slides):
        print(f"--- SLIDE {i+1} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                print(shape.text_frame.text)
else:
    print("PPTX not found!")
