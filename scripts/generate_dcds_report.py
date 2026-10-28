"""
LabTrack: Institutional Hardware Lifecycle, Requisition and Custody Platform
DCDSL Final Academic Project Report Generator
Output: presentations/DCDS/LabTrack_DCDSL_Final_Project_Report.docx

Faithfully built on top of the official faculty template:
presentations/DCDS/DCDSL PROJECT REPORT format.docx

CRITICAL COMPLIANCE RULES:
1. Zero em dashes (\u2014) in the entire document.
2. Official 10-chapter structure strictly preserved.
3. Front matter and official certificates preserved from template.
4. Header: Project Title, Footer: Department of AI&ML and Page Number.
5. Page numbering starts from Chapter 1 (Introduction) at page 1.
6. Authentic database evidence, queries, and UI screenshots embedded.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def sanitize(text: str) -> str:
    """Strictly remove em dashes and markdown artifacts."""
    if not text:
        return ""
    text = text.replace('\u2014', ' - ')
    text = text.replace('\u2013', ' - ')
    text = text.replace('—', ' - ')
    text = text.replace('–', ' - ')
    text = text.replace('**', '')
    text = text.replace('`', '')
    text = text.replace('$', '')
    return text

def set_cell_background(cell, fill_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_chapter_title(doc, number: int, title: str, add_break: bool = True):
    if add_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

    run_ch = p.add_run(f"CHAPTER {number}\n")
    run_ch.font.name = "Times New Roman"
    run_ch.font.size = Pt(13)
    run_ch.font.bold = True
    run_ch.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    run_title = p.add_run(sanitize(title.upper()))
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(2)
    p_div.paragraph_format.space_after = Pt(14)
    run_div = p_div.add_run("_________________________________________________________________________")
    run_div.font.name = "Times New Roman"
    run_div.font.size = Pt(9)
    run_div.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

def add_heading2(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(sanitize(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    return p

def add_heading3(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(sanitize(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_body_paragraph(doc, text: str, bold_prefix: str = None, space_after: float = 6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15

    if bold_prefix:
        r_pre = p.add_run(sanitize(bold_prefix) + " ")
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    r_text = p.add_run(sanitize(text))
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_bullet_point(doc, text: str, bold_prefix: str = None):
    p = doc.add_paragraph(style='List Paragraph')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15

    r_bullet = p.add_run("- ")
    r_bullet.font.name = "Times New Roman"
    r_bullet.font.size = Pt(11)
    r_bullet.font.bold = True
    r_bullet.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    if bold_prefix:
        r_pre = p.add_run(sanitize(bold_prefix) + " ")
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    r_text = p.add_run(sanitize(text))
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_code_snippet(doc, code_str: str, caption: str = None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(sanitize(caption))
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(9.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="16" w:space="0" w:color="1B365D"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05

    lines = code_str.strip().split("\n")
    for i, line in enumerate(lines):
        r = p.add_run(sanitize(line) + ("\n" if i < len(lines)-1 else ""))
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(0)
    p_post.paragraph_format.space_after = Pt(6)

def add_table_data(doc, headers: list, rows: list, col_widths: list = None, caption: str = None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(3)
        p_cap.paragraph_format.keep_with_next = True
        r_cap = p_cap.add_run(sanitize(caption))
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl, color="CBD5E1", sz="4", val="single")

    hdr_cells = tbl.rows[0].cells
    for col_idx, header_text in enumerate(headers):
        cell = hdr_cells[col_idx]
        if col_widths and col_idx < len(col_widths):
            cell.width = Inches(col_widths[col_idx])
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx > 0 and len(header_text) < 14 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

        r = p.add_run(sanitize(header_text))
        r.font.name = "Times New Roman"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for row_idx, row_data in enumerate(rows):
        row_cells = tbl.rows[row_idx + 1].cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            cell = row_cells[col_idx]
            if col_widths and col_idx < len(col_widths):
                cell.width = Inches(col_widths[col_idx])
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            p = cell.paragraphs[0]
            val_str = str(cell_value)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx > 0 and len(val_str) <= 8 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)

            r = p.add_run(sanitize(val_str))
            r.font.name = "Times New Roman"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(0)
    p_post.paragraph_format.space_after = Pt(6)

def add_figure_image(doc, img_rel_path: str, fig_num: str, caption: str, width_in: float = 6.0):
    root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    full_path = os.path.join(root_dir, img_rel_path)

    if not os.path.exists(full_path):
        print(f"Warning: image path does not exist: {full_path}")
        return

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    run_img = p_img.add_run()
    run_img.add_picture(full_path, width=Inches(width_in))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(10)
    p_cap.paragraph_format.keep_with_next = False

    r_num = p_cap.add_run(f"Figure {fig_num}: ")
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(9.5)
    r_num.font.bold = True
    r_num.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    r_cap = p_cap.add_run(sanitize(caption))
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(9.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

def update_front_matter(doc):
    """Update Cover Page, Certificate, and Table of Contents faithfully from template."""
    # Paragraph 10: Project Title
    p10 = doc.paragraphs[10]
    p10.text = ""
    r10 = p10.add_run("LABTRACK: INSTITUTIONAL HARDWARE LIFECYCLE, REQUISITION & CUSTODY MANAGEMENT SYSTEM")
    r10.font.name = "Times New Roman"
    r10.font.size = Pt(14)
    r10.font.bold = True
    r10.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    # Paragraph 14-16: Student Details
    p14 = doc.paragraphs[14]
    p14.text = ""
    r14 = p14.add_run("Student Name\t\tPRN")
    r14.font.name = "Times New Roman"
    r14.font.size = Pt(12)
    r14.font.bold = True

    p15 = doc.paragraphs[15]
    p15.text = ""
    r15 = p15.add_run("Ishaan Rajeev Chand\t25070126076")
    r15.font.name = "Times New Roman"
    r15.font.size = Pt(12)

    p16 = doc.paragraphs[16]
    p16.text = ""  # Clear secondary placeholder

    # Paragraph 20: Guidance
    p20 = doc.paragraphs[20]
    p20.text = ""
    r20 = p20.add_run("Faculty Mentor / Guide")
    r20.font.name = "Times New Roman"
    r20.font.size = Pt(12)

    # Paragraph 24: Academic Year
    p24 = doc.paragraphs[24]
    p24.text = ""
    r24 = p24.add_run("A.Y. 2025-26")
    r24.font.name = "Times New Roman"
    r24.font.size = Pt(12)
    r24.font.bold = True

    # Section 1: Certificate
    p33 = doc.paragraphs[33]
    p33.text = ""
    cert_text = (
        "This is to certify that the DCDSL Project work entitled \"LabTrack: Institutional Hardware Lifecycle, "
        "Requisition and Custody Management System\" is carried out by Ishaan Rajeev Chand (PRN: 25070126076), "
        "in Artificial Intelligence & Machine Learning, Symbiosis International (Deemed University), Pune during "
        "the academic year 2025-2026."
    )
    r33 = p33.add_run(cert_text)
    r33.font.name = "Times New Roman"
    r33.font.size = Pt(11)

    p44 = doc.paragraphs[44]
    p44.text = ""
    r44 = p44.add_run("Name and Signature of the Guide\t\tDr. Shilpa Bade-Gite Head, Department of AI&ML")
    r44.font.name = "Times New Roman"
    r44.font.size = Pt(11)
    r44.font.bold = True

    # Section 2: Table of Contents Table
    toc_table = doc.tables[0]
    set_table_borders(toc_table, color="CBD5E1", sz="4", val="single")

    toc_data = [
        ("Certificate", "ii"),
        ("1. Introduction", "1"),
        ("2. Problem Statement", "3"),
        ("3. System Architecture and Modules", "5"),
        ("4. Functional Requirements", "8"),
        ("5. Entities, Relationships and Attributes", "11"),
        ("6. Relational Schema and Normalization", "15"),
        ("7. Implementation\n    7.1 Database Setup and Multi-Model Infrastructure\n    7.2 Relational Schema DDL Scripts\n    7.3 Stored Procedures, Views, Indexes and ACID Demonstrations\n    7.4 Institutional Master Dataset and Sample Data", "19"),
        ("8. User Interface and Workflows\n    8.1 Screenshots of the User Interface\n    8.2 Step-by-Step Usage Instructions and Validation Results", "29"),
        ("9. Conclusion and Future Scope", "39"),
        ("10. References", "41")
    ]

    # Header Row
    cell_h0 = toc_table.rows[0].cells[0]
    cell_h1 = toc_table.rows[0].cells[1]
    cell_h0.width = Inches(5.2)
    cell_h1.width = Inches(1.3)
    set_cell_background(cell_h0, "1B365D")
    set_cell_background(cell_h1, "1B365D")
    set_cell_margins(cell_h0, top=100, bottom=100, left=120, right=120)
    set_cell_margins(cell_h1, top=100, bottom=100, left=120, right=120)

    p_h0 = cell_h0.paragraphs[0]
    p_h0.text = ""
    r_h0 = p_h0.add_run("Topic / Chapter Title")
    r_h0.font.name = "Times New Roman"
    r_h0.font.size = Pt(10)
    r_h0.font.bold = True
    r_h0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    p_h1 = cell_h1.paragraphs[0]
    p_h1.text = ""
    p_h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_h1 = p_h1.add_run("Page No.")
    r_h1.font.name = "Times New Roman"
    r_h1.font.size = Pt(10)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for i, (topic, page_str) in enumerate(toc_data):
        row = toc_table.rows[i + 1]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(5.2)
        c1.width = Inches(1.3)
        bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_margins(c0, top=60, bottom=60, left=120, right=120)
        set_cell_margins(c1, top=60, bottom=60, left=120, right=120)

        p0 = c0.paragraphs[0]
        p0.text = ""
        r0 = p0.add_run(topic)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(9.5)
        if topic.startswith(("1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.", "10.")):
            r0.font.bold = True

        p1 = c1.paragraphs[0]
        p1.text = ""
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(page_str)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(9.5)
        r1.font.bold = True

    # Remove template instructional notes (paragraphs 53-55)
    for i in range(len(doc.paragraphs)-1, -1, -1):
        p = doc.paragraphs[i]
        t = p.text.strip()
        if 'Kindly submit only one copy' in t or 'Add Header: Project Title' in t or t.startswith('Note'):
            p._p.getparent().remove(p._p)

def setup_body_section(doc):
    """Add Section 3 for the main report body with header, footer, and page numbering starting at 1."""
    new_sect = doc.add_section(WD_SECTION.NEW_PAGE)
    new_sect.page_width = Inches(8.27)
    new_sect.page_height = Inches(11.69)
    new_sect.top_margin = Inches(1.0)
    new_sect.bottom_margin = Inches(1.0)
    new_sect.left_margin = Inches(1.25)
    new_sect.right_margin = Inches(1.0)

    # Unlink header & footer
    new_sect.header.is_linked_to_previous = False
    new_sect.footer.is_linked_to_previous = False

    # Header: Project Title
    hdr_p = new_sect.header.paragraphs[0]
    hdr_p.text = ""
    hdr_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hdr_r = hdr_p.add_run("LabTrack: Institutional Hardware Lifecycle & Custody Platform")
    hdr_r.font.name = "Times New Roman"
    hdr_r.font.size = Pt(8.5)
    hdr_r.font.italic = True
    hdr_r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Footer: Department of AI&ML and Page Number
    ftr_p = new_sect.footer.paragraphs[0]
    ftr_p.text = "Department of AI&ML\t\tPage "
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    ftr_p._p.append(fldSimple)

    for run in ftr_p.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Set page numbering starting at 1
    sectPr = new_sect._sectPr
    pgNumType = sectPr.find(qn('w:pgNumType'))
    if pgNumType is None:
        pgNumType = OxmlElement('w:pgNumType')
        sectPr.append(pgNumType)
    pgNumType.set(qn('w:start'), '1')

def generate_report():
    """Main generation routine."""
    root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    template_path = os.path.join(root_dir, "presentations", "DCDS", "DCDSL PROJECT REPORT format.docx")
    output_path = os.path.join(root_dir, "presentations", "DCDS", "LabTrack_DCDSL_Final_Project_Report.docx")

    print(f"Loading official faculty template: {template_path}")
    doc = docx.Document(template_path)

    print("Updating front matter (Cover Page, Certificate, Table of Contents)...")
    update_front_matter(doc)

    print("Configuring Section 3 for report body (header, footer, page number=1)...")
    setup_body_section(doc)

    from dcds_chapters import (
        build_chapter_1, build_chapter_2, build_chapter_3,
        build_chapter_4, build_chapter_5
    )
    from dcds_chapters_part2 import (
        build_chapter_6, build_chapter_7, build_chapter_8,
        build_chapter_9, build_chapter_10
    )

    print("Writing Chapter 1: Introduction...")
    # For Chapter 1, don't add initial page break since section break already creates new page
    add_chapter_title(doc, 1, "Introduction", add_break=False)
    # The rest of chapter 1 is inside build_chapter_1, but let's call body
    from dcds_chapters import add_heading2, add_body_paragraph, add_bullet_point
    add_heading2(doc, "1.1 Context and Background")
    add_body_paragraph(
        doc,
        "Modern technical education in Artificial Intelligence and Machine Learning (AI&ML), Robotics, "
        "and Embedded Systems requires intensive practical exposure to specialized, high-performance hardware. "
        "At Symbiosis Institute of Technology (SIT Pune), undergraduate and postgraduate students engage in "
        "rigorous laboratory experiments, capstone research projects, and competitive hackathons. These technical "
        "endeavors necessitate continuous access to costly, fragile, and scarce hardware assets. Such components "
        "range from basic prototyping boards and discrete sensors to advanced Edge AI computing platforms including "
        "the NVIDIA Jetson Nano, Jetson Orin series, Intel RealSense depth sensors, digital storage oscilloscopes, "
        "and autonomous robotic kits like the TurtleBot 3."
    )
    add_body_paragraph(
        doc,
        "Historically, engineering departments manage laboratory hardware through localized paper ledgers, "
        "disparate spreadsheets, or manual physical sign-out slips maintained by laboratory attendants. "
        "While viable for small inventories of low-cost passive components, traditional manual approaches break "
        "down catastrophically when applied across modern multi-laboratory academic departments. In an institutional "
        "environment with more than ten distinct laboratory spaces spanning multiple departmental floors, physical "
        "custody tracking becomes fragmented. Students frequently encounter phantom stock, where an online spreadsheet "
        "or verbal confirmation suggests component availability, yet physical units are missing, damaged, or "
        "unaccounted for. Concurrently, faculty mentors face administrative overhead in signing manual authorization slips, "
        "and departmental heads lack real-time visibility into equipment utilization rates and unmet component demand."
    )

    add_heading2(doc, "1.2 Platform Overview: LabTrack")
    add_body_paragraph(
        doc,
        "To resolve these institutional bottlenecks, the LabTrack platform was engineered as an enterprise-grade, "
        "polyglot hardware lifecycle, requisition, and custody governance system. LabTrack replaces manual record-keeping "
        "with an authoritative digital operating environment that bridges student project workflows with faculty governance "
        "and physical counter operations."
    )
    add_body_paragraph(
        doc,
        "The platform introduces end-to-end digital lifecycle governance across every stage of institutional hardware "
        "utilization:",
        bold_prefix="Core Operational Scope:"
    )
    add_bullet_point(
        doc,
        "Centralized multi-lab hardware catalogue categorizing components by engineering tier, functional kind (ASSET vs CONSUMABLE), "
        "and real-time per-lab physical balances.",
        bold_prefix="Catalogue & Inventory Governance:"
    )
    add_bullet_point(
        doc,
        "Formal student research project groupings, cart-based multi-item requisitions, duration bounds, and academic justification.",
        bold_prefix="Requisition Workflows:"
    )
    add_bullet_point(
        doc,
        "Role-based multi-tier approval state machine routing orders through Project Mentors, Lab In-Charges, and Head of Department.",
        bold_prefix="Multi-Tier Approvals:"
    )
    add_bullet_point(
        doc,
        "Strict physical counter verification, barcoded asset binding, digital issuance, and technical return inspection with condition grading.",
        bold_prefix="Custody Handover & Return:"
    )
    add_bullet_point(
        doc,
        "Asynchronous projection of out-of-stock student requisitions to quantify unfulfilled hardware demand for departmental budgeting.",
        bold_prefix="Unmet Demand Analytics:"
    )

    add_heading2(doc, "1.3 Multi-Model Polyglot Architecture")
    add_body_paragraph(
        doc,
        "A foundational achievement of the LabTrack system is its dual-database polyglot persistence architecture. "
        "Recognizing that distinct institutional data models require divergent storage guarantees, LabTrack combines an "
        "authoritative relational database with an analytical document store:"
    )
    add_bullet_point(
        doc,
        "Aiven Cloud MySQL 8.4 Enterprise serves as the authoritative source of truth for transactional state, "
        "institutional identity, role-based permissions, financial custody, row-level inventory locks, and immutable compliance audits. "
        "It enforces strict ACID semantics, preventing double-allocation and ensuring complete data integrity.",
        bold_prefix="Authoritative Relational Engine (Aiven MySQL 8.4):"
    )
    add_bullet_point(
        doc,
        "MongoDB Atlas v7.0 operates alongside the relational core to manage polymorphic component datasheets, "
        "real-time student event notification streams, and asynchronous unmet demand projections. It utilizes multi-stage "
        "aggregation pipelines to summarize procurement signals without placing analytical load on the transactional relational core.",
        bold_prefix="Analytical Document Store (MongoDB Atlas v7.0):"
    )

    add_heading2(doc, "1.4 Scope and Objectives")
    add_body_paragraph(
        doc,
        "The objective of this academic project is to design, implement, and validate a production-ready institutional "
        "hardware management platform fulfilling the requirements of the Database Concepts for Data Science Lab (DCDS) "
        "curriculum at Symbiosis Institute of Technology. The platform satisfies all requirements of conceptual ER modeling, "
        "relational schema normalization up to Boyce-Codd Normal Form (BCNF), physical MySQL DDL implementation, complex relational "
        "JOINs, subqueries, stored procedures, database views, execution plan optimization, and multi-model document aggregation."
    )

    print("Writing Chapter 2: Problem Statement...")
    build_chapter_2(doc)

    print("Writing Chapter 3: System Architecture and Modules...")
    build_chapter_3(doc)

    print("Writing Chapter 4: Functional Requirements...")
    build_chapter_4(doc)

    print("Writing Chapter 5: Entities, Relationships and Attributes...")
    build_chapter_5(doc)

    print("Writing Chapter 6: Relational Schema and Normalization...")
    build_chapter_6(doc)

    print("Writing Chapter 7: Implementation (Database Setup, DDL, SQL Evidence, ACID)...")
    build_chapter_7(doc)

    print("Writing Chapter 8: User Interface (Screenshots and Workflows)...")
    build_chapter_8(doc)

    print("Writing Chapter 9: Conclusion and Future Scope...")
    build_chapter_9(doc)

    print("Writing Chapter 10: References...")
    build_chapter_10(doc)

    print(f"Saving final report to: {output_path}")
    doc.save(output_path)

    # Verification Pass
    print("\nExecuting comprehensive verification pass...")
    assert os.path.exists(output_path), "Output report file does not exist!"
    file_size_kb = os.path.getsize(output_path) / 1024
    print(f"Report file generated successfully ({file_size_kb:.1f} KB).")

    # Em dash scanner
    verify_doc = docx.Document(output_path)
    em_dash_violations = []

    for p_idx, p in enumerate(verify_doc.paragraphs):
        if '\u2014' in p.text or '—' in p.text:
            em_dash_violations.append(f"Paragraph {p_idx}: {p.text[:60]}")

    for t_idx, table in enumerate(verify_doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                if '\u2014' in cell.text or '—' in cell.text:
                    em_dash_violations.append(f"Table {t_idx} [R{r_idx}C{c_idx}]: {cell.text[:60]}")

    for s_idx, s in enumerate(verify_doc.sections):
        for p in s.header.paragraphs:
            if '\u2014' in p.text or '—' in p.text:
                em_dash_violations.append(f"Header in Section {s_idx}: {p.text}")
        for p in s.footer.paragraphs:
            if '\u2014' in p.text or '—' in p.text:
                em_dash_violations.append(f"Footer in Section {s_idx}: {p.text}")

    if em_dash_violations:
        print(f"ERROR: Found {len(em_dash_violations)} EM DASH violations!")
        for v in em_dash_violations[:10]:
            print("  ", v)
        raise ValueError(f"Strict compliance failed: {len(em_dash_violations)} em dashes detected.")

    print("COMPLIANCE VERIFIED: ZERO em dashes (\\u2014) in the entire document!")
    print(f"Total Sections: {len(verify_doc.sections)}")
    print(f"Total Paragraphs: {len(verify_doc.paragraphs)}")
    print(f"Total Tables: {len(verify_doc.tables)}")

if __name__ == "__main__":
    generate_report()
