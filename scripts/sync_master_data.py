"""Authoritative Master Data Synchronization Tool for LabTrack.
Synchronizes data between public/data/labtrack_master_data.xlsx, Relational SQL
(SQLite / MySQL), and NoSQL Document/GridFS (MongoDB Atlas).
"""
import argparse
import hashlib
import mimetypes
import os
import sys
from datetime import datetime, timezone
import openpyxl
from dotenv import load_dotenv

# Ensure backend package can be imported
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "src", "backend")))

from labtrack.config import settings
from labtrack import db
from labtrack.models import Base, User, Lab, Item, Stock, Asset, Project, Member, utcnow
from labtrack.auth import password_hash

load_dotenv()

ALT_EXCEL_PATH = os.path.join("public", "data", "labtrack_master_data_updated.xlsx")
EXCEL_PATH = ALT_EXCEL_PATH if os.path.exists(ALT_EXCEL_PATH) else os.path.join("public", "data", "labtrack_master_data.xlsx")
HARDWARE_IMG_DIR = os.path.join("public", "images", "hardware")


def load_workbook_data(path=EXCEL_PATH):
    if not os.path.exists(path):
        raise FileNotFoundError(f"Master workbook not found at: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
        data = [dict(zip(headers, row)) for row in rows[1:] if any(row)]
        sheets[name] = data
    return sheets


def ensure_table_columns(session):
    """Ensure any newly added model columns are migrated into existing database tables."""
    from sqlalchemy import text
    dialect = db.engine.dialect.name
    if dialect == "sqlite":
        needed_cols = {
            "users": [
                ("department", "VARCHAR(50)"),
                ("designation", "VARCHAR(100)"),
                ("batch", "INTEGER"),
                ("year_of_study", "INTEGER"),
                ("branch", "VARCHAR(20)"),
                ("roll_number", "VARCHAR(20)"),
                ("phone", "VARCHAR(30)"),
            ],
            "labs": [
                ("lab_code", "VARCHAR(30)"),
                ("floor", "INTEGER"),
                ("wing", "VARCHAR(20)"),
                ("location_desc", "VARCHAR(200)"),
                ("branch", "VARCHAR(20)"),
            ],
            "items": [
                ("image_url", "VARCHAR(255)"),
                ("cost_tier", "VARCHAR(50)"),
            ]
        }
        with db.engine.connect() as conn:
            for table, cols in needed_cols.items():
                existing_cols = [r[1] for r in conn.execute(text(f"PRAGMA table_info({table})")).fetchall()]
                for col_name, col_type in cols:
                    if col_name not in existing_cols:
                        conn.execute(text(f"ALTER TABLE {table} ADD COLUMN {col_name} {col_type}"))
                        print(f"  [+] Migrated column: {table}.{col_name} ({col_type})")
            conn.commit()


def sync_relational_database(sheets, demo_password=None):
    demo_password = demo_password or os.getenv("LABTRACK_DEMO_PASSWORD", "LabTrackDemo2026!")
    hashed_pw = password_hash(demo_password)
    print(f"[*] Synchronizing Relational DB ({db.engine.dialect.name})...")

    Base.metadata.create_all(db.engine)
    ensure_table_columns(db.Session())

    with db.Session() as session:
        # 1. Labs Pre-Mapping (need Lab IDs for Assistants, and User IDs for Incharges)
        labs_data = sheets.get("Labs", [])
        users_data = sheets.get("Users", [])

        # Create Users
        user_identifier_to_id = {}
        for idx, u in enumerate(users_data, start=1):
            identifier = str(u["User ID"]).strip()
            name = str(u["Full Name"]).strip()
            email = str(u["Email"]).strip()
            user_type = str(u["User Type"]).strip().upper()
            dept = str(u.get("Department") or "CENTRAL").strip()
            desig = str(u.get("Designation") or "").strip()
            batch = int(u["Batch"]) if u.get("Batch") not in [None, ""] else None
            year = int(u["Year of Study"]) if u.get("Year of Study") not in [None, ""] else None
            branch = str(u.get("Branch") or "").strip() or None
            roll = str(u.get("Roll Number") or "").strip() or None
            phone = str(u.get("Phone") or "").strip()

            # Assign roles
            roles_list = []
            if user_type == "STUDENT":
                roles_list = ["STUDENT"]
            elif user_type == "HOD":
                roles_list = ["HOD", "FACULTY"]
            elif user_type == "ADMIN":
                roles_list = ["ADMIN"]
            elif user_type == "ASSISTANT":
                roles_list = ["ASSISTANT"]
            else: # FACULTY
                roles_list = ["FACULTY", "MENTOR"]
                if u.get("In-Charge Labs"):
                    roles_list.append("INCHARGE")

            existing = session.query(User).filter((User.identifier == identifier) | (User.email == email)).first()
            if existing:
                existing.name = name
                existing.identifier = identifier
                existing.email = email
                existing.roles = roles_list
                existing.department = dept
                existing.designation = desig
                existing.batch = batch
                existing.year_of_study = year
                existing.branch = branch
                existing.roll_number = roll
                existing.phone = phone
                existing.password_hash = hashed_pw
                session.flush()
                user_identifier_to_id[identifier] = existing.id
            else:
                new_user = User(
                    name=name,
                    identifier=identifier,
                    email=email,
                    password_hash=hashed_pw,
                    roles=roles_list,
                    lab_ids=[],
                    active=True,
                    department=dept,
                    designation=desig,
                    batch=batch,
                    year_of_study=year,
                    branch=branch,
                    roll_number=roll,
                    phone=phone
                )
                session.add(new_user)
                session.flush()
                user_identifier_to_id[identifier] = new_user.id

        # Create Labs
        lab_code_to_id = {}
        for idx, l in enumerate(labs_data, start=1):
            lab_code = str(l["Lab ID"]).strip()
            name = str(l["Lab Name"]).strip()
            floor = int(l["Floor"])
            wing = str(l.get("Wing") or "")
            loc = str(l["Location Description"]).strip()
            branch = str(l["Branch"]).strip()
            incharge_uid_str = str(l["In-Charge Faculty ID"]).strip()
            incharge_id = user_identifier_to_id.get(incharge_uid_str, 1)

            existing_lab = session.query(Lab).filter((Lab.lab_code == lab_code) | (Lab.name == name)).first()
            if existing_lab:
                existing_lab.name = name
                existing_lab.lab_code = lab_code
                existing_lab.incharge_id = incharge_id
                existing_lab.floor = floor
                existing_lab.wing = wing
                existing_lab.location_desc = loc
                existing_lab.branch = branch
                session.flush()
                lab_code_to_id[lab_code] = existing_lab.id
            else:
                new_lab = Lab(
                    name=name,
                    incharge_id=incharge_id,
                    lab_code=lab_code,
                    floor=floor,
                    wing=wing,
                    location_desc=loc,
                    branch=branch
                )
                session.add(new_lab)
                session.flush()
                lab_code_to_id[lab_code] = new_lab.id

        # Update lab_ids on Users (Assistants & Incharges)
        for u in users_data:
            identifier = str(u["User ID"]).strip()
            uid = user_identifier_to_id.get(identifier)
            if not uid:
                continue
            user_obj = session.get(User, uid)
            assigned_lab_ids = []
            if u.get("Managed Labs"):
                for code in str(u["Managed Labs"]).split(","):
                    code = code.strip()
                    if code in lab_code_to_id:
                        assigned_lab_ids.append(lab_code_to_id[code])
            if u.get("In-Charge Labs"):
                for code in str(u["In-Charge Labs"]).split(","):
                    code = code.strip()
                    if code in lab_code_to_id:
                        assigned_lab_ids.append(lab_code_to_id[code])
            user_obj.lab_ids = assigned_lab_ids
        session.flush()

        # 3. Create Items
        items_data = sheets.get("Items", [])
        item_code_to_id = {}
        for idx, itm in enumerate(items_data, start=1):
            item_code = str(itm["Item ID"]).strip()
            name = str(itm["Item Name"]).strip()
            kind = str(itm["Kind"]).strip().upper()
            category = str(itm["Category"]).strip()
            brand = str(itm.get("Manufacturer / Brand") or "").strip()
            desc = str(itm.get("Description") or "").strip()
            specs_str = str(itm.get("Sub-Specifications") or "").strip()
            cost_tier = str(itm.get("Cost Tier") or "").strip()
            image_url = str(itm.get("Image URL") or "").strip()

            specs_dict = {"specs": specs_str, "brand": brand, "category": category}

            existing_item = session.query(Item).filter(Item.name == name).first()
            if existing_item:
                existing_item.name = name
                existing_item.kind = kind
                existing_item.category = category
                existing_item.brand = brand
                existing_item.description = desc
                existing_item.specs = specs_dict
                existing_item.cost_tier = cost_tier
                existing_item.image_url = image_url
                session.flush()
                item_code_to_id[item_code] = existing_item.id
            else:
                new_item = Item(
                    name=name,
                    kind=kind,
                    category=category,
                    description=desc,
                    brand=brand,
                    specs=specs_dict,
                    alternative_ids=[],
                    cost_tier=cost_tier,
                    image_url=image_url
                )
                session.add(new_item)
                session.flush()
                item_code_to_id[item_code] = new_item.id

        # 4. Create Stock
        stock_data = sheets.get("Stock", [])
        for s_row in stock_data:
            lab_code = str(s_row["Lab ID"]).strip()
            item_code = str(s_row["Item ID"]).strip()
            good = int(s_row["Good Quantity"])
            bad = int(s_row.get("Bad Quantity", 0))

            lab_id = lab_code_to_id.get(lab_code)
            item_id = item_code_to_id.get(item_code)
            if not lab_id or not item_id:
                continue

            existing_stock = session.query(Stock).filter(Stock.lab_id == lab_id, Stock.item_id == item_id).first()
            if existing_stock:
                existing_stock.good = good
                existing_stock.bad = bad
            else:
                session.add(Stock(lab_id=lab_id, item_id=item_id, good=good, bad=bad))

        # 5. Create Assets
        assets_data = sheets.get("Assets", [])
        for a_row in assets_data:
            barcode = str(a_row["Barcode"]).strip()
            item_code = str(a_row["Item ID"]).strip()
            lab_code = str(a_row["Lab ID"]).strip()
            default_id = str(a_row["Default ID"]).strip()
            bill = str(a_row["Bill Number"]).strip()
            state = str(a_row.get("State", "AVAILABLE")).strip().upper()

            item_id = item_code_to_id.get(item_code)
            lab_id = lab_code_to_id.get(lab_code)
            if not item_id or not lab_id:
                continue

            existing_asset = session.query(Asset).filter(Asset.barcode == barcode).first()
            if existing_asset:
                existing_asset.item_id = item_id
                existing_asset.lab_id = lab_id
                existing_asset.default_id = default_id
                existing_asset.bill_number = bill
                existing_asset.state = state
            else:
                session.add(Asset(barcode=barcode, item_id=item_id, lab_id=lab_id, default_id=default_id, bill_number=bill, state=state))

        # 6. Create Projects & Members
        projects_data = sheets.get("Projects", [])
        members_data = sheets.get("Project_Members", [])
        proj_code_to_id = {}

        for p_row in projects_data:
            code = str(p_row["Project ID"]).strip()
            name = str(p_row["Project Name"]).strip()
            desc = str(p_row.get("Description") or "").strip()
            lead_prn = str(p_row["Lead PRN"]).strip()
            mentor_id_str = str(p_row["Mentor Faculty ID"]).strip()

            lead_uid = user_identifier_to_id.get(lead_prn, 1)
            mentor_uid = user_identifier_to_id.get(mentor_id_str, 1)

            existing_proj = session.query(Project).filter(Project.name == name).first()
            if existing_proj:
                existing_proj.name = name
                existing_proj.description = desc
                existing_proj.lead_id = lead_uid
                existing_proj.mentor_id = mentor_uid
                session.flush()
                proj_code_to_id[code] = existing_proj.id
            else:
                new_proj = Project(name=name, description=desc, lead_id=lead_uid, mentor_id=mentor_uid)
                session.add(new_proj)
                session.flush()
                proj_code_to_id[code] = new_proj.id

        for m_row in members_data:
            p_code = str(m_row["Project ID"]).strip()
            s_prn = str(m_row["Student PRN"]).strip()
            p_id = proj_code_to_id.get(p_code)
            s_id = user_identifier_to_id.get(s_prn)
            if not p_id or not s_id:
                continue
            if not session.get(Member, (p_id, s_id)):
                session.add(Member(project_id=p_id, user_id=s_id))

        session.commit()
    print("[OK] Relational Database synchronization completed successfully!")


def sync_mongodb_and_images(sheets):
    mongo_url = os.getenv("LABTRACK_MONGO_URL")
    if not mongo_url:
        print("[SKIP] LABTRACK_MONGO_URL not set; skipping MongoDB synchronization.")
        return

    print("[*] Connecting to MongoDB Atlas...")
    try:
        from pymongo import MongoClient
        from gridfs import GridFS
        client = MongoClient(mongo_url, serverSelectionTimeoutMS=5000)
        mongo_db = client[os.getenv("LABTRACK_MONGO_DATABASE", "labtrack")]
        fs = GridFS(mongo_db)

        # 1. Store Component Documents in MongoDB for DCDS NoSQL evaluation
        comp_col = mongo_db["component_catalog_documents"]
        items_data = sheets.get("Items", [])
        for itm in items_data:
            item_id = str(itm["Item ID"]).strip()
            doc = {
                "item_id": item_id,
                "name": str(itm["Item Name"]).strip(),
                "kind": str(itm["Kind"]).strip().upper(),
                "category": str(itm["Category"]).strip(),
                "brand": str(itm.get("Manufacturer / Brand") or "").strip(),
                "sub_specs": str(itm.get("Sub-Specifications") or "").strip(),
                "description": str(itm.get("Description") or "").strip(),
                "cost_tier": str(itm.get("Cost Tier") or "").strip(),
                "image_url": str(itm.get("Image URL") or "").strip(),
                "last_synchronized": datetime.now(timezone.utc)
            }
            comp_col.update_one({"item_id": item_id}, {"$set": doc}, upsert=True)
        print(f"[OK] Synchronized {len(items_data)} items into MongoDB collection 'component_catalog_documents'")

        # 2. Store Binary Images in MongoDB GridFS and document collection 'component_images'
        img_col = mongo_db["component_images"]
        if os.path.exists(HARDWARE_IMG_DIR):
            for fname in os.listdir(HARDWARE_IMG_DIR):
                if not fname.lower().endswith((".jpg", ".jpeg", ".png", ".webp")):
                    continue
                fpath = os.path.join(HARDWARE_IMG_DIR, fname)
                with open(fpath, "rb") as f:
                    content = f.read()
                sha256 = hashlib.sha256(content).hexdigest()
                mime_type = mimetypes.guess_type(fname)[0] or "image/jpeg"

                # Check GridFS
                existing_file = fs.find_one({"filename": fname})
                if existing_file:
                    gridfs_id = existing_file._id
                else:
                    gridfs_id = fs.put(content, filename=fname, content_type=mime_type, sha256=sha256)

                img_doc = {
                    "filename": fname,
                    "sha256": sha256,
                    "mime_type": mime_type,
                    "size_bytes": len(content),
                    "gridfs_id": str(gridfs_id),
                    "cdn_url": f"/images/hardware/{fname}",
                    "uploaded_at": datetime.now(timezone.utc)
                }
                img_col.update_one({"filename": fname}, {"$set": img_doc}, upsert=True)
            print(f"[OK] Ingested hardware images into MongoDB GridFS and collection 'component_images'")
    except Exception as e:
        print(f"[WARN] MongoDB Atlas sync error: {e}")


def main():
    parser = argparse.ArgumentParser(description="LabTrack Authoritative Master Data Sync Tool")
    parser.add_argument("--import-all", action="store_true", help="Sync Excel data into Relational DB and MongoDB Atlas")
    parser.add_argument("--validate-only", action="store_true", help="Validate Excel workbook structure without mutating DB")
    args = parser.parse_args()

    sheets = load_workbook_data()
    print(f"Loaded {len(sheets)} sheets from {EXCEL_PATH}: {list(sheets.keys())}")
    for name, data in sheets.items():
        print(f"  - Sheet '{name}': {len(data)} records")

    if args.validate_only:
        print("[OK] Workbook validation passed cleanly.")
        return

    sync_relational_database(sheets)
    sync_mongodb_and_images(sheets)
    print("\n[SUCCESS] Master data and component image synchronization finished!")


if __name__ == "__main__":
    main()
