"""
LabTrack: Institutional Hardware Lifecycle, Requisition and Custody Platform
DCDSL Final Academic Project Report Refinement Generator
Output: presentations/DCDS/LabTrack_DCDSL_Final_Project_Report_Refined.docx

Faithfully built on top of the official faculty template:
presentations/DCDS/DCDSL PROJECT REPORT format.docx

CRITICAL COMPLIANCE RULES:
1. Zero em dashes (\u2014) in the entire document.
2. Official 10-chapter structure strictly preserved.
3. Front matter and official certificates preserved from template.
4. Header: Project Title, Footer: Department of AI&ML and Page Number.
5. Page numbering starts from Chapter 1 (Introduction) at page 1.
6. 6 professional academic diagrams embedded.
7. Completely restructured relational schema domain tables & attribute dictionaries.
8. Professional UI screenshot placeholders prepared for Raghu.
9. Word proofing underlines disabled on code/identifiers (<w:noProof/>).
10. Multipage table repeated headers (<w:tblHeader/>) and row split prevention (<w:cantSplit/>).
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def sanitize(text: str) -> str:
    """Strictly remove em dashes and markdown artifacts."""
    if not text:
        return ""
    text = text.replace('\u2014', ' - ')
    text = text.replace('\u2013', ' - ')
    text = text.replace('—', ' - ')
    text = text.replace('–', ' - ')
    text = text.replace('**', '')
    text = text.replace('`', '')
    text = text.replace('$', '')
    return text

def disable_proofing(run):
    """Disable spell check and grammar check underlines in Word."""
    rPr = run._r.get_or_add_rPr()
    rPr.append(parse_xml(f'<w:noProof {nsdecls("w")}/>'))

def set_cell_background(cell, fill_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="8", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_chapter_title(doc, number: int, title: str, add_break: bool = True):
    if add_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

    run_ch = p.add_run(f"CHAPTER {number}\n")
    run_ch.font.name = "Times New Roman"
    run_ch.font.size = Pt(13)
    run_ch.font.bold = True
    run_ch.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    run_title = p.add_run(sanitize(title.upper()))
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(15.5)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(2)
    p_div.paragraph_format.space_after = Pt(12)
    run_div = p_div.add_run("_________________________________________________________________________")
    run_div.font.name = "Times New Roman"
    run_div.font.size = Pt(9)
    run_div.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

def add_heading2(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(sanitize(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    return p

def add_heading3(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(sanitize(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_body_paragraph(doc, text: str, bold_prefix: str = None, space_after: float = 6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15

    if bold_prefix:
        r_pre = p.add_run(sanitize(bold_prefix) + " ")
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    r_text = p.add_run(sanitize(text))
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_bullet_point(doc, text: str, bold_prefix: str = None):
    p = doc.add_paragraph(style='List Paragraph')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15

    r_bullet = p.add_run("- ")
    r_bullet.font.name = "Times New Roman"
    r_bullet.font.size = Pt(11)
    r_bullet.font.bold = True
    r_bullet.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    if bold_prefix:
        r_pre = p.add_run(sanitize(bold_prefix) + " ")
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    r_text = p.add_run(sanitize(text))
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_code_snippet(doc, code_str: str, caption: str = None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        p_cap.paragraph_format.keep_with_next = True
        r_cap = p_cap.add_run(sanitize(caption))
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(9.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "FFFFFF")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="16" w:space="0" w:color="000000"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05

    lines = code_str.strip().split("\n")
    for i, line in enumerate(lines):
        r = p.add_run(sanitize(line) + ("\n" if i < len(lines)-1 else ""))
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        disable_proofing(r)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(0)
    p_post.paragraph_format.space_after = Pt(6)

def add_table_data(doc, headers: list, rows: list, col_widths: list = None, caption: str = None, right_align_cols: list = None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(3)
        p_cap.paragraph_format.keep_with_next = True
        r_cap = p_cap.add_run(sanitize(caption))
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(9.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl, color="CBD5E1", sz="4", val="single")

    # Header Row - Pure White Background with Bold Black Text
    hdr_row = tbl.rows[0]
    trPr0 = hdr_row._tr.get_or_add_trPr()
    trPr0.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr0.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    for col_idx, header_text in enumerate(headers):
        cell = hdr_row.cells[col_idx]
        if col_widths and col_idx < len(col_widths):
            cell.width = Inches(col_widths[col_idx])
        set_cell_background(cell, "FFFFFF")
        set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # Bold border under header row
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if right_align_cols and col_idx in right_align_cols else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

        r = p.add_run(sanitize(header_text))
        r.font.name = "Times New Roman"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # Data Rows - Pure White Background with Black Text
    for row_idx, row_data in enumerate(rows):
        row = tbl.rows[row_idx + 1]
        trPr_r = row._tr.get_or_add_trPr()
        trPr_r.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        bg_color = "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            cell = row.cells[col_idx]
            if col_widths and col_idx < len(col_widths):
                cell.width = Inches(col_widths[col_idx])
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            # Subtle horizontal divider on data rows, bottom row has single black border
            tcPr = cell._tc.get_or_add_tcPr()
            b_bottom = "000000" if row_idx == len(rows)-1 else "E2E8F0"
            b_sz = "8" if row_idx == len(rows)-1 else "4"
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:top w:val="none"/>'
                f'<w:bottom w:val="single" w:sz="{b_sz}" w:space="0" w:color="{b_bottom}"/>'
                f'<w:left w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(borders)

            p = cell.paragraphs[0]
            val_str = str(cell_value)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if right_align_cols and col_idx in right_align_cols else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)

            r = p.add_run(sanitize(val_str))
            r.font.name = "Times New Roman"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            # Disable proofing for technical symbols/identifiers
            if any(ch in val_str for ch in ['_', '(', ')', '->', 'VARCHAR', 'INT', 'BOOLEAN', 'DATETIME', 'CHECK', 'UNIQUE']):
                disable_proofing(r)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(0)
    p_post.paragraph_format.space_after = Pt(6)

def add_figure_image(doc, img_rel_path: str, fig_num: str, caption: str, width_in: float = 6.2):
    root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    full_path = os.path.join(root_dir, img_rel_path)

    if not os.path.exists(full_path):
        print(f"Warning: image path does not exist: {full_path}")
        return

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.paragraph_format.keep_with_next = True
    run_img = p_img.add_run()
    run_img.add_picture(full_path, width=Inches(width_in))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(10)
    p_cap.paragraph_format.keep_with_next = False

    r_num = p_cap.add_run(f"Figure {fig_num}: ")
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(9.5)
    r_num.font.bold = True
    r_num.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    r_cap = p_cap.add_run(sanitize(caption))
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(9.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

def add_ui_placeholder(doc, fig_num: str, title: str, route: str, dimensions: str, purpose: str, artifact: str):
    """Render a dedicated, professional placeholder frame for Raghu's upcoming frontend screenshots."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "FFFFFF")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:left w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:right w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15

    r_tag = p.add_run(f"[UI SCREENSHOT PLACEHOLDER: FIGURE {fig_num}]\n")
    r_tag.font.name = "Times New Roman"
    r_tag.font.size = Pt(10)
    r_tag.font.bold = True
    r_tag.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    r_meta = p.add_run(f"Target Route: {route} | Viewport: {dimensions}\n")
    r_meta.font.name = "Consolas"
    r_meta.font.size = Pt(8.5)
    r_meta.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    disable_proofing(r_meta)

    r_art = p.add_run(f"Artifact Source: {artifact}\n\n")
    r_art.font.name = "Times New Roman"
    r_art.font.size = Pt(8.5)
    r_art.font.italic = True
    r_art.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    r_purp = p.add_run(f"Verification Objective: {purpose}")
    r_purp.font.name = "Times New Roman"
    r_purp.font.size = Pt(9)
    r_purp.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # Caption below placeholder
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(12)
    p_cap.paragraph_format.keep_with_next = False

    r_num = p_cap.add_run(f"Figure {fig_num}: ")
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(9.5)
    r_num.font.bold = True
    r_num.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    r_t = p_cap.add_run(sanitize(title))
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(9.5)
    r_t.font.italic = True
    r_t.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

def update_front_matter(doc):
    """Update Cover Page, Certificate, and Table of Contents faithfully from template."""
    # Paragraph 10: Project Title
    p10 = doc.paragraphs[10]
    p10.text = ""
    r10 = p10.add_run("LABTRACK: INSTITUTIONAL HARDWARE LIFECYCLE, REQUISITION & CUSTODY MANAGEMENT SYSTEM")
    r10.font.name = "Times New Roman"
    r10.font.size = Pt(14)
    r10.font.bold = True
    r10.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # Paragraph 14-16: Student Details
    p14 = doc.paragraphs[14]
    p14.text = ""
    r14 = p14.add_run("Student Name\t\tPRN")
    r14.font.name = "Times New Roman"
    r14.font.size = Pt(12)
    r14.font.bold = True

    p15 = doc.paragraphs[15]
    p15.text = ""
    r15 = p15.add_run("Ishaan Rajeev Chand\t25070126076")
    r15.font.name = "Times New Roman"
    r15.font.size = Pt(12)

    p16 = doc.paragraphs[16]
    p16.text = ""

    # Paragraph 20: Guidance
    p20 = doc.paragraphs[20]
    p20.text = ""
    r20 = p20.add_run("Faculty Mentor / Guide")
    r20.font.name = "Times New Roman"
    r20.font.size = Pt(12)

    # Paragraph 24: Academic Year
    p24 = doc.paragraphs[24]
    p24.text = ""
    r24 = p24.add_run("A.Y. 2025-26")
    r24.font.name = "Times New Roman"
    r24.font.size = Pt(12)
    r24.font.bold = True

    # Section 1: Certificate
    p33 = doc.paragraphs[33]
    p33.text = ""
    cert_text = (
        "This is to certify that the DCDSL Project work entitled \"LabTrack: Institutional Hardware Lifecycle, "
        "Requisition and Custody Management System\" is carried out by Ishaan Rajeev Chand (PRN: 25070126076), "
        "in Artificial Intelligence & Machine Learning, Symbiosis International (Deemed University), Pune during "
        "the academic year 2025-2026."
    )
    r33 = p33.add_run(cert_text)
    r33.font.name = "Times New Roman"
    r33.font.size = Pt(11)

    p44 = doc.paragraphs[44]
    p44.text = ""
    r44 = p44.add_run("Name and Signature of the Guide\t\tDr. Shilpa Bade-Gite Head, Department of AI&ML")
    r44.font.name = "Times New Roman"
    r44.font.size = Pt(11)
    r44.font.bold = True

    # Section 2: Table of Contents Table
    toc_table = doc.tables[0]
    set_table_borders(toc_table, color="CBD5E1", sz="4", val="single")

    toc_data = [
        ("Certificate", "ii"),
        ("1. Introduction", "1"),
        ("2. Problem Statement", "3"),
        ("3. System Architecture and Modules", "5"),
        ("4. Functional Requirements", "7"),
        ("5. Entities, Relationships and Attributes", "10"),
        ("6. Relational Schema and Normalization\n    6.1 Schema Overview and Architectural Division\n    6.2 Relational Schema by Functional Domain\n    6.3 Attribute-Level Data Dictionary\n    6.4 Key and Constraint Verification Audit\n    6.5 Functional Dependencies\n    6.6 Normalization Analysis (1NF to BCNF)", "13"),
        ("7. Implementation\n    7.1 Database Setup and Multi-Model Infrastructure\n    7.2 Table Creation Scripts\n    7.3 Stored Procedures, Views, Indexes and ACID Demonstrations\n    7.4 Institutional Master Dataset and Sample Data", "20"),
        ("8. User Interface and Workflows\n    8.1 Screenshots of the User Interface (Insertion Placeholders)\n    8.2 Step-by-Step Usage Instructions and Role Workflows\n    8.3 Automated Verification and Quality Assurance", "28"),
        ("9. Conclusion and Future Scope", "32"),
        ("10. References", "34")
    ]

    cell_h0 = toc_table.rows[0].cells[0]
    cell_h1 = toc_table.rows[0].cells[1]
    cell_h0.width = Inches(5.2)
    cell_h1.width = Inches(1.3)
    set_cell_background(cell_h0, "FFFFFF")
    set_cell_background(cell_h1, "FFFFFF")
    set_cell_margins(cell_h0, top=90, bottom=90, left=120, right=120)
    set_cell_margins(cell_h1, top=90, bottom=90, left=120, right=120)

    for c in [cell_h0, cell_h1]:
        tcPr = c._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

    p_h0 = cell_h0.paragraphs[0]
    p_h0.text = ""
    r_h0 = p_h0.add_run("Topic / Chapter Title")
    r_h0.font.name = "Times New Roman"
    r_h0.font.size = Pt(10)
    r_h0.font.bold = True
    r_h0.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    p_h1 = cell_h1.paragraphs[0]
    p_h1.text = ""
    p_h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_h1 = p_h1.add_run("Page No.")
    r_h1.font.name = "Times New Roman"
    r_h1.font.size = Pt(10)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    for i, (topic, page_str) in enumerate(toc_data):
        row = toc_table.rows[i + 1]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(5.2)
        c1.width = Inches(1.3)
        bg = "FFFFFF"
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_margins(c0, top=60, bottom=60, left=120, right=120)
        set_cell_margins(c1, top=60, bottom=60, left=120, right=120)

        p0 = c0.paragraphs[0]
        p0.text = ""
        r0 = p0.add_run(topic)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        if topic.startswith(("1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.", "10.")):
            r0.font.bold = True

        p1 = c1.paragraphs[0]
        p1.text = ""
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(page_str)
        r1.font.name = "Times New Roman"
        r1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        r1.font.size = Pt(9.5)
        r1.font.bold = True

    # Remove template instruction notes
    for i in range(len(doc.paragraphs)-1, -1, -1):
        p = doc.paragraphs[i]
        t = p.text.strip()
        if 'Kindly submit only one copy' in t or 'Add Header: Project Title' in t or t.startswith('Note'):
            p._p.getparent().remove(p._p)

def setup_body_section(doc):
    """Add Section 3 for report body with header, footer, and page numbering starting at 1."""
    new_sect = doc.add_section(WD_SECTION.NEW_PAGE)
    new_sect.page_width = Inches(8.27)
    new_sect.page_height = Inches(11.69)
    new_sect.top_margin = Inches(1.0)
    new_sect.bottom_margin = Inches(1.0)
    new_sect.left_margin = Inches(1.25)
    new_sect.right_margin = Inches(1.0)
    new_sect.header_distance = Inches(0.5)
    new_sect.footer_distance = Inches(0.5)

    new_sect.header.is_linked_to_previous = False
    new_sect.footer.is_linked_to_previous = False

    TAB_STOP_RIGHT = "8666"

    # Header: Symmetric two-column with bottom border rule
    hdr = new_sect.header
    for p in hdr.paragraphs:
        p._p.getparent().remove(p._p)

    hdr_p_xml = f'''
    <w:p {nsdecls("w")}>
      <w:pPr>
        <w:pBdr>
          <w:bottom w:val="single" w:sz="6" w:space="4" w:color="94A3B8"/>
        </w:pBdr>
        <w:tabs>
          <w:tab w:val="right" w:pos="{TAB_STOP_RIGHT}"/>
        </w:tabs>
        <w:spacing w:before="0" w:after="72" w:line="240" w:lineRule="auto"/>
      </w:pPr>
      <w:r>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
          <w:b/>
          <w:color w:val="000000"/>
          <w:sz w:val="19"/>
          <w:szCs w:val="19"/>
          <w:noProof/>
        </w:rPr>
        <w:t>LabTrack: Institutional Hardware Lifecycle &amp; Custody System</w:t>
      </w:r>
      <w:r>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
          <w:i/>
          <w:color w:val="475569"/>
          <w:sz w:val="19"/>
          <w:szCs w:val="19"/>
          <w:noProof/>
        </w:rPr>
        <w:tab/>
        <w:t>DCDSL Project Report</w:t>
      </w:r>
    </w:p>
    '''
    hdr._element.append(parse_xml(hdr_p_xml))

    # Footer: Symmetric two-column with top border rule
    ftr = new_sect.footer
    for p in ftr.paragraphs:
        p._p.getparent().remove(p._p)

    ftr_p_xml = f'''
    <w:p {nsdecls("w")}>
      <w:pPr>
        <w:pBdr>
          <w:top w:val="single" w:sz="6" w:space="4" w:color="94A3B8"/>
        </w:pBdr>
        <w:tabs>
          <w:tab w:val="right" w:pos="{TAB_STOP_RIGHT}"/>
        </w:tabs>
        <w:spacing w:before="72" w:after="0" w:line="240" w:lineRule="auto"/>
      </w:pPr>
      <w:r>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
          <w:b/>
          <w:color w:val="000000"/>
          <w:sz w:val="19"/>
          <w:szCs w:val="19"/>
          <w:noProof/>
        </w:rPr>
        <w:t>Department of AI&amp;ML</w:t>
      </w:r>
      <w:r>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
          <w:b/>
          <w:color w:val="1E293B"/>
          <w:sz w:val="19"/>
          <w:szCs w:val="19"/>
          <w:noProof/>
        </w:rPr>
        <w:tab/>
        <w:t xml:space="preserve">Page </w:t>
      </w:r>
      <w:fldSimple w:instr="PAGE">
        <w:r>
          <w:rPr>
            <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
            <w:b/>
            <w:color w:val="1E293B"/>
            <w:sz w:val="19"/>
            <w:szCs w:val="19"/>
            <w:noProof/>
          </w:rPr>
          <w:t>1</w:t>
        </w:r>
      </w:fldSimple>
    </w:p>
    '''
    ftr._element.append(parse_xml(ftr_p_xml))

    # Page number start at 1
    sectPr = new_sect._sectPr
    pgNumType = sectPr.find(qn('w:pgNumType'))
    if pgNumType is None:
        pgNumType = OxmlElement('w:pgNumType')
        sectPr.append(pgNumType)
    pgNumType.set(qn('w:start'), '1')

def generate_refined_report():
    root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    template_path = os.path.join(root_dir, "presentations", "DCDS", "DCDSL PROJECT REPORT format.docx")
    output_path = os.path.join(root_dir, "presentations", "DCDS", "LabTrack_DCDSL_Final_Project_Report_Refined.docx")

    print(f"Loading official faculty template: {template_path}")
    doc = docx.Document(template_path)

    print("Updating front matter (Cover Page, Certificate, Table of Contents)...")
    update_front_matter(doc)

    print("Configuring Section 3 for report body (header, footer, page number=1)...")
    setup_body_section(doc)

    from refined_chapters_part1 import (
        build_chapter_1, build_chapter_2, build_chapter_3,
        build_chapter_4, build_chapter_5
    )
    from refined_chapters_part2 import (
        build_chapter_6, build_chapter_7, build_chapter_8,
        build_chapter_9, build_chapter_10
    )

    print("Writing Chapter 1: Introduction...")
    build_chapter_1(doc)

    print("Writing Chapter 2: Problem Statement...")
    build_chapter_2(doc)

    print("Writing Chapter 3: System Architecture and Modules (Figure 1)...")
    build_chapter_3(doc)

    print("Writing Chapter 4: Functional Requirements (Figure 2)...")
    build_chapter_4(doc)

    print("Writing Chapter 5: Entities, Relationships and Attributes (Figure 3)...")
    build_chapter_5(doc)

    print("Writing Chapter 6: Relational Schema and Normalization (Restructured, Figure 4)...")
    build_chapter_6(doc)

    print("Writing Chapter 7: Implementation (Figures 5 and 6, DDL, SQL, ACID)...")
    build_chapter_7(doc)

    print("Writing Chapter 8: User Interface (Placeholders 8.1 - 8.12, Instructions)...")
    build_chapter_8(doc)

    print("Writing Chapter 9: Conclusion and Future Scope...")
    build_chapter_9(doc)

    print("Writing Chapter 10: References...")
    build_chapter_10(doc)

    print(f"Saving refined report to: {output_path}")
    doc.save(output_path)

    # Verification Pass
    print("\nExecuting comprehensive verification pass...")
    assert os.path.exists(output_path), "Refined report file does not exist!"
    file_size_kb = os.path.getsize(output_path) / 1024
    print(f"Report file generated successfully ({file_size_kb:.1f} KB).")

    verify_doc = docx.Document(output_path)
    em_dash_violations = []

    for p_idx, p in enumerate(verify_doc.paragraphs):
        if '\u2014' in p.text or '—' in p.text:
            em_dash_violations.append(f"Paragraph {p_idx}: {p.text[:60]}")

    for t_idx, table in enumerate(verify_doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                if '\u2014' in cell.text or '—' in cell.text:
                    em_dash_violations.append(f"Table {t_idx} [R{r_idx}C{c_idx}]: {cell.text[:60]}")

    for s_idx, s in enumerate(verify_doc.sections):
        for p in s.header.paragraphs:
            if '\u2014' in p.text or '—' in p.text:
                em_dash_violations.append(f"Header in Section {s_idx}: {p.text}")
        for p in s.footer.paragraphs:
            if '\u2014' in p.text or '—' in p.text:
                em_dash_violations.append(f"Footer in Section {s_idx}: {p.text}")

    if em_dash_violations:
        print(f"ERROR: Found {len(em_dash_violations)} EM DASH violations!")
        for v in em_dash_violations[:10]:
            print("  ", v)
        raise ValueError(f"Strict compliance failed: {len(em_dash_violations)} em dashes detected.")

    print("COMPLIANCE VERIFIED: ZERO em dashes (\\u2014) in the entire document!")
    print(f"Total Sections: {len(verify_doc.sections)}")
    print(f"Total Paragraphs: {len(verify_doc.paragraphs)}")
    print(f"Total Tables: {len(verify_doc.tables)}")

if __name__ == "__main__":
    generate_refined_report()
